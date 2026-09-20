import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse

from .models import Division, Department, CostCenter, WorkLocation
from attendance.models import CompanyHoliday


def _require_super_admin(user):
    return user.is_authenticated and user.role == 'admin'


# ─── Overview ────────────────────────────────────────────────
@login_required
def masterdata_overview(request):
    if not _require_super_admin(request.user):
        messages.error(request, 'เฉพาะ Super Admin เท่านั้น')
        return redirect('dashboard')

    tab = request.GET.get('tab', 'division')

    # Holiday year filter
    from datetime import date as date_cls
    current_year = date_cls.today().year
    try:
        holiday_year = int(request.GET.get('holiday_year', current_year))
    except (ValueError, TypeError):
        holiday_year = current_year

    all_holiday_years = (
        CompanyHoliday.objects
        .values_list('date__year', flat=True)
        .distinct()
        .order_by('date__year')
    )
    holiday_years = sorted(set(list(all_holiday_years) + [current_year]))

    context = {
        'tab': tab,
        'divisions': Division.objects.all().order_by('code'),
        'departments': Department.objects.select_related('division').order_by('code'),
        'cost_centers': CostCenter.objects.all().order_by('code'),
        'work_locations': WorkLocation.objects.all().order_by('name'),
        'holidays': CompanyHoliday.objects.filter(date__year=holiday_year).order_by('date'),
        'holiday_year': holiday_year,
        'holiday_years': holiday_years,
    }
    return render(request, 'employees/masterdata/overview.html', context)


# ─── Division ────────────────────────────────────────────────
@login_required
def division_create(request):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    if request.method == 'POST':
        code = request.POST.get('code', '').strip()
        name = request.POST.get('name', '').strip()
        if not code or not name:
            messages.error(request, 'กรุณากรอก Code และ Name')
        elif Division.objects.filter(code=code).exists():
            messages.error(request, f'Code "{code}" มีอยู่แล้ว')
        else:
            Division.objects.create(code=code, name=name)
            messages.success(request, f'เพิ่ม Division "{name}" เรียบร้อย')
    return redirect('masterdata_overview')


@login_required
def division_edit(request, pk):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    obj = get_object_or_404(Division, pk=pk)
    if request.method == 'POST':
        obj.code = request.POST.get('code', '').strip()
        obj.name = request.POST.get('name', '').strip()
        obj.save()
        messages.success(request, 'แก้ไข Division เรียบร้อย')
        return redirect('masterdata_overview')
    return render(request, 'employees/masterdata/edit_form.html', {
        'obj': obj, 'type': 'division', 'fields': [
            {'name': 'code', 'label': 'Code', 'value': obj.code},
            {'name': 'name', 'label': 'Name', 'value': obj.name},
        ]
    })


@login_required
def division_delete(request, pk):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    obj = get_object_or_404(Division, pk=pk)
    if request.method == 'POST':
        obj.delete()
        messages.success(request, 'ลบ Division เรียบร้อย')
    return redirect('masterdata_overview')


# ─── Department ──────────────────────────────────────────────
@login_required
def department_create(request):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    if request.method == 'POST':
        code = request.POST.get('code', '').strip()
        name = request.POST.get('name', '').strip()
        division_id = request.POST.get('division') or None
        if not code or not name:
            messages.error(request, 'กรุณากรอก Code และ Name')
        elif Department.objects.filter(code=code).exists():
            messages.error(request, f'Code "{code}" มีอยู่แล้ว')
        else:
            Department.objects.create(code=code, name=name, division_id=division_id)
            messages.success(request, f'เพิ่ม Department "{name}" เรียบร้อย')
    return redirect('/employees/masterdata/?tab=department')


@login_required
def department_edit(request, pk):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    obj = get_object_or_404(Department, pk=pk)
    if request.method == 'POST':
        obj.code = request.POST.get('code', '').strip()
        obj.name = request.POST.get('name', '').strip()
        obj.division_id = request.POST.get('division') or None
        obj.save()
        messages.success(request, 'แก้ไข Department เรียบร้อย')
        return redirect('masterdata_overview')
    return render(request, 'employees/masterdata/edit_form.html', {
        'obj': obj, 'type': 'department',
        'divisions': Division.objects.all().order_by('code'),
        'fields': [
            {'name': 'code', 'label': 'Code', 'value': obj.code},
            {'name': 'name', 'label': 'Name', 'value': obj.name},
        ]
    })


@login_required
def department_delete(request, pk):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    obj = get_object_or_404(Department, pk=pk)
    if request.method == 'POST':
        obj.delete()
        messages.success(request, 'ลบ Department เรียบร้อย')
    return redirect('masterdata_overview')


# ─── Cost Center ─────────────────────────────────────────────
@login_required
def costcenter_create(request):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    if request.method == 'POST':
        code = request.POST.get('code', '').strip()
        name = request.POST.get('name', '').strip()
        if not code or not name:
            messages.error(request, 'กรุณากรอก Code และ Name')
        elif CostCenter.objects.filter(code=code).exists():
            messages.error(request, f'Code "{code}" มีอยู่แล้ว')
        else:
            CostCenter.objects.create(code=code, name=name)
            messages.success(request, f'เพิ่ม Cost Center "{name}" เรียบร้อย')
    return redirect('/employees/masterdata/?tab=costcenter')


@login_required
def costcenter_edit(request, pk):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    obj = get_object_or_404(CostCenter, pk=pk)
    if request.method == 'POST':
        obj.code = request.POST.get('code', '').strip()
        obj.name = request.POST.get('name', '').strip()
        obj.save()
        messages.success(request, 'แก้ไข Cost Center เรียบร้อย')
        return redirect('masterdata_overview')
    return render(request, 'employees/masterdata/edit_form.html', {
        'obj': obj, 'type': 'costcenter', 'fields': [
            {'name': 'code', 'label': 'Code', 'value': obj.code},
            {'name': 'name', 'label': 'Name', 'value': obj.name},
        ]
    })


@login_required
def costcenter_delete(request, pk):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    obj = get_object_or_404(CostCenter, pk=pk)
    if request.method == 'POST':
        obj.delete()
        messages.success(request, 'ลบ Cost Center เรียบร้อย')
    return redirect('masterdata_overview')


# ─── Work Location ───────────────────────────────────────────
@login_required
def worklocation_create(request):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        address = request.POST.get('address', '').strip()
        if not name:
            messages.error(request, 'กรุณากรอก Name')
        else:
            WorkLocation.objects.create(name=name, address=address)
            messages.success(request, f'เพิ่ม Work Location "{name}" เรียบร้อย')
    return redirect('/employees/masterdata/?tab=worklocation')


@login_required
def worklocation_edit(request, pk):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    obj = get_object_or_404(WorkLocation, pk=pk)
    if request.method == 'POST':
        obj.name = request.POST.get('name', '').strip()
        obj.address = request.POST.get('address', '').strip()
        obj.save()
        messages.success(request, 'แก้ไข Work Location เรียบร้อย')
        return redirect('masterdata_overview')
    return render(request, 'employees/masterdata/edit_form.html', {
        'obj': obj, 'type': 'worklocation', 'fields': [
            {'name': 'name', 'label': 'Name', 'value': obj.name},
            {'name': 'address', 'label': 'Address', 'value': obj.address},
        ]
    })


@login_required
def worklocation_delete(request, pk):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    obj = get_object_or_404(WorkLocation, pk=pk)
    if request.method == 'POST':
        obj.delete()
        messages.success(request, 'ลบ Work Location เรียบร้อย')
    return redirect('masterdata_overview')


# ─── Excel Template Download ─────────────────────────────────
@login_required
def masterdata_excel_template(request):
    if not _require_super_admin(request.user):
        return redirect('dashboard')

    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='152057')

    def make_sheet(wb, title, headers, examples):
        ws = wb.create_sheet(title=title)
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            ws.column_dimensions[cell.column_letter].width = 24
        for row_idx, row in enumerate(examples, 2):
            for col_idx, val in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)
        return ws

    # Remove default sheet
    wb.remove(wb.active)

    make_sheet(wb, 'Division',
        ['code', 'name'],
        [['DIV001', 'Operations'], ['DIV002', 'Finance']])

    make_sheet(wb, 'Department',
        ['code', 'name', 'division_code'],
        [['DEPT001', 'Accounting', 'DIV002'], ['DEPT002', 'Logistics', 'DIV001']])

    make_sheet(wb, 'CostCenter',
        ['code', 'name'],
        [['CC001', 'Head Office'], ['CC002', 'Warehouse']])

    make_sheet(wb, 'WorkLocation',
        ['name', 'address'],
        [['Bangkok HQ', '123 Sathorn Rd, Bangkok'],
         ['Samut Prakan', '456 Factory Rd, SP']])

    # Instructions sheet
    ws_info = wb.create_sheet(title='Instructions', index=0)
    ws_info['A1'] = 'HRMS Master Data Upload Template'
    ws_info['A1'].font = Font(bold=True, size=14)
    instructions = [
        ('', ''),
        ('Sheet', 'Description'),
        ('Division', 'code (unique), name'),
        ('Department', 'code (unique), name, division_code (must match Division.code)'),
        ('CostCenter', 'code (unique), name'),
        ('WorkLocation', 'name, address (optional)'),
        ('', ''),
        ('Note', 'Row 1 = headers (do not change). Data starts row 2.'),
        ('Note', 'Existing records with same code will be UPDATED (not duplicated).'),
    ]
    for r, (a, b) in enumerate(instructions, 2):
        ws_info.cell(row=r, column=1, value=a).font = Font(bold=True) if a == 'Sheet' else Font()
        ws_info.cell(row=r, column=2, value=b)
    ws_info.column_dimensions['A'].width = 20
    ws_info.column_dimensions['B'].width = 60

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="masterdata_template.xlsx"'
    return response


# ─── Excel Upload ─────────────────────────────────────────────
@login_required
def masterdata_excel_upload(request):
    if not _require_super_admin(request.user):
        return redirect('dashboard')

    if request.method != 'POST':
        return redirect('masterdata_overview')

    file = request.FILES.get('excel_file')
    if not file:
        messages.error(request, 'กรุณาเลือกไฟล์ Excel')
        return redirect('masterdata_overview')

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
    except Exception:
        messages.error(request, 'ไฟล์ไม่ถูกต้อง กรุณาใช้ไฟล์ .xlsx')
        return redirect('masterdata_overview')

    results = []
    errors = []

    # Division
    if 'Division' in wb.sheetnames:
        ws = wb['Division']
        added = updated = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            code, name = (str(row[0]).strip() if row[0] else ''), (str(row[1]).strip() if row[1] else '')
            if not code or not name:
                continue
            obj, created = Division.objects.update_or_create(code=code, defaults={'name': name})
            if created: added += 1
            else: updated += 1
        results.append(f'Division: เพิ่ม {added}, อัปเดต {updated}')

    # Department
    if 'Department' in wb.sheetnames:
        ws = wb['Department']
        added = updated = skipped = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = str(row[0]).strip() if row[0] else ''
            name = str(row[1]).strip() if row[1] else ''
            div_code = str(row[2]).strip() if len(row) > 2 and row[2] else ''
            if not code or not name:
                continue
            division = Division.objects.filter(code=div_code).first() if div_code else None
            obj, created = Department.objects.update_or_create(
                code=code, defaults={'name': name, 'division': division}
            )
            if created: added += 1
            else: updated += 1
        results.append(f'Department: เพิ่ม {added}, อัปเดต {updated}')

    # CostCenter
    if 'CostCenter' in wb.sheetnames:
        ws = wb['CostCenter']
        added = updated = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = str(row[0]).strip() if row[0] else ''
            name = str(row[1]).strip() if row[1] else ''
            if not code or not name:
                continue
            obj, created = CostCenter.objects.update_or_create(code=code, defaults={'name': name})
            if created: added += 1
            else: updated += 1
        results.append(f'Cost Center: เพิ่ม {added}, อัปเดต {updated}')

    # WorkLocation
    if 'WorkLocation' in wb.sheetnames:
        ws = wb['WorkLocation']
        added = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = str(row[0]).strip() if row[0] else ''
            address = str(row[1]).strip() if len(row) > 1 and row[1] else ''
            if not name:
                continue
            WorkLocation.objects.get_or_create(name=name, defaults={'address': address})
            added += 1
        results.append(f'Work Location: นำเข้า {added}')

    if results:
        messages.success(request, 'นำเข้าข้อมูลสำเร็จ — ' + ' | '.join(results))
    if errors:
        for e in errors:
            messages.error(request, e)

    return redirect('masterdata_overview')


# ─── Company Holiday ─────────────────────────────────────────
@login_required
def holiday_create(request):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    if request.method == 'POST':
        from datetime import date as date_cls
        date_str = request.POST.get('date', '').strip()
        name = request.POST.get('name', '').strip()
        name_en = request.POST.get('name_en', '').strip()
        if not date_str or not name:
            messages.error(request, 'กรุณากรอกวันที่และชื่อวันหยุด')
        else:
            try:
                from datetime import datetime
                d = datetime.strptime(date_str, '%Y-%m-%d').date()
                obj, created = CompanyHoliday.objects.update_or_create(
                    date=d, defaults={'name': name, 'name_en': name_en}
                )
                if created:
                    messages.success(request, f'เพิ่มวันหยุด "{name}" เรียบร้อย')
                else:
                    messages.success(request, f'อัปเดตวันหยุด "{name}" เรียบร้อย')
            except ValueError:
                messages.error(request, 'รูปแบบวันที่ไม่ถูกต้อง')
    year = request.POST.get('date', '')[:4] or ''
    redirect_url = f'/employees/masterdata/?tab=holiday&holiday_year={year}' if year else '/employees/masterdata/?tab=holiday'
    return redirect(redirect_url)


@login_required
def holiday_delete(request, pk):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    obj = get_object_or_404(CompanyHoliday, pk=pk)
    year = obj.date.year
    if request.method == 'POST':
        obj.delete()
        messages.success(request, 'ลบวันหยุดเรียบร้อย')
    return redirect(f'/employees/masterdata/?tab=holiday&holiday_year={year}')


@login_required
def holiday_excel_template(request):
    if not _require_super_admin(request.user):
        return redirect('dashboard')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Holidays'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='152057')
    headers = ['date', 'name', 'name_en']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 22

    # Example rows
    examples = [
        ['2026-01-01', 'วันขึ้นปีใหม่', "New Year's Day"],
        ['2026-04-13', 'วันสงกรานต์', 'Songkran Festival'],
        ['2026-05-01', 'วันแรงงานแห่งชาติ', 'Labour Day'],
        ['2026-12-31', 'วันสิ้นปี', "New Year's Eve"],
    ]
    for r_idx, row in enumerate(examples, 2):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="holiday_template.xlsx"'
    return response


@login_required
def holiday_excel_upload(request):
    if not _require_super_admin(request.user):
        return redirect('dashboard')
    if request.method != 'POST':
        return redirect('masterdata_overview')

    file = request.FILES.get('excel_file')
    if not file:
        messages.error(request, 'กรุณาเลือกไฟล์ Excel')
        return redirect('/employees/masterdata/?tab=holiday')

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
    except Exception:
        messages.error(request, 'ไฟล์ไม่ถูกต้อง กรุณาใช้ไฟล์ .xlsx')
        return redirect('/employees/masterdata/?tab=holiday')

    ws = wb.active
    added = updated = skipped = 0
    from datetime import datetime, date as date_cls

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        raw_date = row[0]
        name = str(row[1]).strip() if row[1] else ''
        name_en = str(row[2]).strip() if len(row) > 2 and row[2] else ''
        if not name:
            skipped += 1
            continue
        # Parse date: support datetime obj (from Excel) or string
        if isinstance(raw_date, (datetime, date_cls)):
            d = raw_date.date() if isinstance(raw_date, datetime) else raw_date
        else:
            try:
                d = datetime.strptime(str(raw_date).strip(), '%Y-%m-%d').date()
            except ValueError:
                skipped += 1
                continue
        obj, created = CompanyHoliday.objects.update_or_create(
            date=d, defaults={'name': name, 'name_en': name_en}
        )
        if created:
            added += 1
        else:
            updated += 1

    msg = f'นำเข้าวันหยุด: เพิ่ม {added}, อัปเดต {updated}'
    if skipped:
        msg += f', ข้ามแถว {skipped}'
    messages.success(request, msg)
    return redirect('/employees/masterdata/?tab=holiday')


# ═══════════════════════════════════════════════════════════════
#  EMPLOYEE IMPORT / EXPORT
# ═══════════════════════════════════════════════════════════════

EMPLOYEE_HEADERS = [
    'employee_id', 'username', 'password',
    'title_prefix', 'first_name_en', 'last_name_en',
    'first_name_th', 'last_name_th', 'nickname',
    'gender', 'national_id', 'date_of_birth',
    'company_email', 'personal_email', 'phone',
    'employee_status', 'employment_type', 'join_date',
    'job_title', 'job_level',
    'department_code', 'division_code', 'cost_center_code', 'work_location_name',
    'direct_manager_employee_id',
]


@login_required
def employee_excel_template(request):
    """Download blank template (with 1 example row)."""
    if not _require_super_admin(request.user):
        return redirect('dashboard')

    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='152057')
    note_font = Font(color='FF0000', italic=True)

    ws = wb.active
    ws.title = 'Employees'

    for col, h in enumerate(EMPLOYEE_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 26

    # Example row
    example = [
        'EMP001', 'john.doe', 'Password123!',
        'mr', 'John', 'Doe',
        'จอห์น', 'โด', 'Johnny',
        'male', '1234567890123', '1990-01-15',
        'john.doe@company.com', 'john@gmail.com', '081-234-5678',
        'active', 'full_time', '2024-01-01',
        'Software Engineer', 'P3',
        'DEPT001', 'DIV-CORP', 'CC001', 'Bangkok HQ',
        '',
    ]
    for col, val in enumerate(example, 1):
        ws.cell(row=2, column=col, value=val)

    # Instructions sheet
    ws2 = wb.create_sheet('Instructions')
    rows = [
        ('HRMS Employee Import Template', ''),
        ('', ''),
        ('Column', 'Notes'),
        ('employee_id', 'Required. Unique key — used to update existing employees.'),
        ('username', 'Required for NEW employees. Leave blank to skip creating a login.'),
        ('password', 'Required for NEW employees. Ignored when updating existing.'),
        ('title_prefix', 'mr / mrs / ms / dr / other'),
        ('first_name_en', 'Required.'),
        ('last_name_en', 'Required.'),
        ('gender', 'male / female / other'),
        ('date_of_birth', 'YYYY-MM-DD'),
        ('employee_status', 'active / probation / resigned / terminated / retired'),
        ('employment_type', 'full_time / part_time / contract / intern / probation'),
        ('join_date', 'YYYY-MM-DD'),
        ('department_code', 'Must match Department code in Master Data'),
        ('division_code', 'Must match Division code in Master Data'),
        ('cost_center_code', 'Must match Cost Center code in Master Data'),
        ('work_location_name', 'Must match Work Location name in Master Data'),
        ('direct_manager_employee_id', 'employee_id of the direct manager (must exist)'),
        ('', ''),
        ('Note', 'Rows with existing employee_id are UPDATED; new employee_id rows are CREATED.'),
        ('Note', 'Row 1 = headers (do not change). Data starts at row 2.'),
    ]
    for r, (a, b) in enumerate(rows, 1):
        ca = ws2.cell(row=r, column=1, value=a)
        cb = ws2.cell(row=r, column=2, value=b)
        if r in (1, 3):
            ca.font = Font(bold=True, size=(13 if r == 1 else 11))
            cb.font = Font(bold=True)
    ws2.column_dimensions['A'].width = 32
    ws2.column_dimensions['B'].width = 65

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="employee_import_template.xlsx"'
    return response


@login_required
def employee_excel_export(request):
    """Export all existing employees to Excel (same column layout as template)."""
    if not _require_super_admin(request.user):
        return redirect('dashboard')

    from .models import EmployeeProfile
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='152057')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Employees'

    for col, h in enumerate(EMPLOYEE_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 26

    profiles = EmployeeProfile.objects.select_related(
        'user', 'department', 'division', 'cost_center', 'work_location', 'direct_manager'
    ).order_by('employee_id')

    for row_idx, p in enumerate(profiles, 2):
        dob = p.date_of_birth.strftime('%Y-%m-%d') if p.date_of_birth else ''
        join = p.join_date.strftime('%Y-%m-%d') if p.join_date else ''
        values = [
            p.employee_id,
            p.user.username if p.user else '',
            '',  # password — never export
            p.title_prefix,
            p.first_name_en, p.last_name_en,
            p.first_name_th, p.last_name_th,
            p.nickname, p.gender,
            p.national_id, dob,
            p.company_email, p.personal_email, p.phone,
            p.employee_status, p.employment_type, join,
            p.job_title, p.job_level,
            p.department.code if p.department else '',
            p.division.code if p.division else '',
            p.cost_center.code if p.cost_center else '',
            p.work_location.name if p.work_location else '',
            p.direct_manager.employee_id if p.direct_manager else '',
        ]
        for col_idx, val in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="employees_export.xlsx"'
    return response


@login_required
def employee_excel_upload(request):
    """Import employees from uploaded Excel file."""
    if not _require_super_admin(request.user):
        return redirect('dashboard')

    if request.method != 'POST':
        return redirect('employee_list')

    from django.contrib.auth import get_user_model
    from datetime import datetime
    from .models import EmployeeProfile

    User = get_user_model()

    file = request.FILES.get('excel_file')
    if not file:
        messages.error(request, 'กรุณาเลือกไฟล์ Excel')
        return redirect('employee_list')

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
    except Exception:
        messages.error(request, 'ไฟล์ไม่ถูกต้อง กรุณาใช้ไฟล์ .xlsx')
        return redirect('employee_list')

    if 'Employees' not in wb.sheetnames:
        messages.error(request, 'ไม่พบ sheet "Employees" ในไฟล์')
        return redirect('employee_list')

    ws = wb['Employees']
    headers = [str(c.value).strip() if c.value else '' for c in ws[1]]

    def col(name):
        try:
            return headers.index(name)
        except ValueError:
            return None

    def cell(row, name):
        idx = col(name)
        if idx is None or idx >= len(row):
            return ''
        v = row[idx]
        return str(v).strip() if v is not None else ''

    def parse_date(s):
        if not s:
            return None
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
        return None

    created_count = updated_count = skipped_count = 0
    errors = []
    # Two-pass: collect manager assignments
    manager_map = {}  # employee_id -> direct_manager_employee_id

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for row_num, row in enumerate(rows, 2):
        if all(v is None for v in row):
            continue

        emp_id = cell(row, 'employee_id')
        first_en = cell(row, 'first_name_en')
        last_en = cell(row, 'last_name_en')

        if not emp_id:
            skipped_count += 1
            continue
        if not first_en or not last_en:
            errors.append(f'Row {row_num}: employee_id={emp_id} — ต้องมี first_name_en และ last_name_en')
            skipped_count += 1
            continue

        # Resolve FK lookups
        dept_code = cell(row, 'department_code')
        div_code = cell(row, 'division_code')
        cc_code = cell(row, 'cost_center_code')
        wl_name = cell(row, 'work_location_name')

        department = Department.objects.filter(code=dept_code).first() if dept_code else None
        division = Division.objects.filter(code=div_code).first() if div_code else None
        cost_center = CostCenter.objects.filter(code=cc_code).first() if cc_code else None
        work_location = WorkLocation.objects.filter(name=wl_name).first() if wl_name else None

        # User account
        username = cell(row, 'username')
        password = cell(row, 'password')
        user = None

        existing_profile = EmployeeProfile.objects.filter(employee_id=emp_id).first()
        if existing_profile:
            user = existing_profile.user
        elif username:
            user, u_created = User.objects.get_or_create(
                username=username,
                defaults={'email': cell(row, 'company_email') or f'{username}@company.com'}
            )
            if u_created and password:
                user.set_password(password)
                user.role = 'employee'
                user.save()

        defaults = {
            'first_name_en': first_en,
            'last_name_en': last_en,
            'title_prefix': cell(row, 'title_prefix'),
            'first_name_th': cell(row, 'first_name_th'),
            'last_name_th': cell(row, 'last_name_th'),
            'nickname': cell(row, 'nickname'),
            'gender': cell(row, 'gender'),
            'national_id': cell(row, 'national_id'),
            'date_of_birth': parse_date(cell(row, 'date_of_birth')),
            'company_email': cell(row, 'company_email'),
            'personal_email': cell(row, 'personal_email'),
            'phone': cell(row, 'phone'),
            'employee_status': cell(row, 'employee_status') or 'active',
            'employment_type': cell(row, 'employment_type') or 'full_time',
            'join_date': parse_date(cell(row, 'join_date')),
            'job_title': cell(row, 'job_title'),
            'job_level': cell(row, 'job_level'),
            'department': department,
            'division': division,
            'cost_center': cost_center,
            'work_location': work_location,
        }
        if user:
            defaults['user'] = user

        try:
            profile, p_created = EmployeeProfile.objects.update_or_create(
                employee_id=emp_id, defaults=defaults
            )
            if p_created:
                created_count += 1
            else:
                updated_count += 1

            mgr_id = cell(row, 'direct_manager_employee_id')
            if mgr_id:
                manager_map[emp_id] = mgr_id

        except Exception as e:
            errors.append(f'Row {row_num}: employee_id={emp_id} — {e}')
            skipped_count += 1

    # Second pass: assign managers
    for emp_id, mgr_id in manager_map.items():
        mgr = EmployeeProfile.objects.filter(employee_id=mgr_id).first()
        if mgr:
            EmployeeProfile.objects.filter(employee_id=emp_id).update(direct_manager=mgr)

    summary = f'สร้างใหม่ {created_count} | อัปเดต {updated_count} | ข้าม {skipped_count}'
    messages.success(request, f'นำเข้า Employee เสร็จสิ้น — {summary}')
    for e in errors[:5]:
        messages.warning(request, e)
    if len(errors) > 5:
        messages.warning(request, f'...และอีก {len(errors) - 5} error')

    return redirect('employee_list')
